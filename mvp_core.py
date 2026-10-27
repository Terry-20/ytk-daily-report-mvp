"""MVP 概念验证计算核心：货款 = 产品数量 × 单价。

同时被 Web 应用 (app.py) 与 CLI (scripts/mvp_calc.py) 复用，
保证两条路径计算结果一致。
"""

import openpyxl


def _first_number_in_column(ws, header_keyword):
    """在表中定位含 header_keyword 的表头列，返回该列第一个有效数字。

    找不到表头或无有效数字时抛出中文 ValueError。
    """
    max_col = ws.max_column or 0
    header_col = None
    for col in range(1, max_col + 1):
        val = ws.cell(row=1, column=col).value
        if val is not None and header_keyword in str(val):
            header_col = col
            break
    if header_col is None:
        raise ValueError(f"表头未找到「{header_keyword}」")

    max_row = ws.max_row or 0
    for row in range(2, max_row + 1):
        v = ws.cell(row=row, column=header_col).value
        if v is None:
            continue
        try:
            num = float(v)
        except (TypeError, ValueError):
            continue
        return num
    raise ValueError(f"「{header_keyword}」列无有效数字")


def compute_payment(material1_path, material2_path):
    """读取材料1/材料2，返回 货款 = 产品数量 × 单价。"""
    try:
        wb1 = openpyxl.load_workbook(material1_path, data_only=True)
        ws1 = wb1.active
    except Exception as e:
        raise ValueError(f"读取材料1失败：{e}")
    try:
        wb2 = openpyxl.load_workbook(material2_path, data_only=True)
        ws2 = wb2.active
    except Exception as e:
        raise ValueError(f"读取材料2失败：{e}")

    qty = _first_number_in_column(ws1, "产品数量")
    price = _first_number_in_column(ws2, "单价")
    return qty * price


def build_result(out_path, value):
    """生成成品 结果.xlsx：单 sheet，表头仅「货款」，下一行写数值。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "结果"
    ws.append(["货款"])
    ws.append([value])
    wb.save(out_path)
